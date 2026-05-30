from __future__ import annotations

import argparse
import json
import os
import sys
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
RUNTIME_PYTHON_PACKAGES = Path(
    r"C:\Users\Sai\.cache\codex-runtimes\codex-primary-runtime\dependencies\python"
)
if RUNTIME_PYTHON_PACKAGES.exists():
    sys.path.insert(0, str(RUNTIME_PYTHON_PACKAGES))

from openpyxl import load_workbook  # type: ignore


def chunked(values: list[Any], size: int) -> list[list[Any]]:
    return [values[index : index + size] for index in range(0, len(values), size)]


def load_env_file(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def resolve_env(key: str) -> str | None:
    if os.environ.get(key):
        return os.environ[key]
    for candidate in (ROOT / ".env", ROOT / ".env.local"):
        data = load_env_file(candidate)
        if key in data:
            return data[key]
    return None


@dataclass
class SupabaseRest:
    base_url: str
    service_role_key: str

    @property
    def headers(self) -> dict[str, str]:
        return {
            "apikey": self.service_role_key,
            "Authorization": f"Bearer {self.service_role_key}",
            "Content-Type": "application/json",
        }

    def request(
        self,
        method: str,
        table: str,
        *,
        query: dict[str, str] | None = None,
        body: Any | None = None,
        extra_headers: dict[str, str] | None = None,
    ) -> Any:
        encoded_query = urllib.parse.urlencode(query or {})
        url = f"{self.base_url}/rest/v1/{table}"
        if encoded_query:
            url = f"{url}?{encoded_query}"

        headers = dict(self.headers)
        if extra_headers:
            headers.update(extra_headers)

        payload = None
        if body is not None:
            payload = json.dumps(body).encode("utf-8")

        request = urllib.request.Request(url, data=payload, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request) as response:
                raw = response.read().decode("utf-8")
                if not raw:
                    return None
                return json.loads(raw)
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"{method} {url} failed: {exc.code} {detail}") from exc


def normalize_phone(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) == 10:
        return f"+91{digits}"
    if len(digits) == 12 and digits.startswith("91"):
        return f"+{digits}"
    if len(digits) > 0 and not digits.startswith("+"):
        return f"+{digits}"
    return digits


def parse_tag_values(value: Any) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    normalized = text.replace("\n", ",").replace(";", ",").replace("|", ",")
    tags: list[str] = []
    seen: set[str] = set()
    for raw in normalized.split(","):
        tag = raw.strip()
        if not tag:
            continue
        key = tag.casefold()
        if key in seen:
            continue
        seen.add(key)
        tags.append(tag)
    return tags


def read_contacts(xlsx_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []

    columns = {str(name).strip().lower(): idx for idx, name in enumerate(header) if name}
    required = {"name", "phone"}
    missing = required - set(columns)
    if missing:
        raise RuntimeError(f"Missing required columns in workbook: {', '.join(sorted(missing))}")

    contacts: list[dict[str, Any]] = []
    seen_phones: set[str] = set()
    for row in rows:
        name = row[columns["name"]] if columns.get("name") is not None else None
        phone = row[columns["phone"]] if columns.get("phone") is not None else None
        email = row[columns["email"]] if "email" in columns else None
        tag_value = None
        if "tags" in columns:
            tag_value = row[columns["tags"]]
        elif "tag" in columns:
            tag_value = row[columns["tag"]]

        normalized_phone = normalize_phone(phone)
        if not normalized_phone or normalized_phone in seen_phones:
            continue
        seen_phones.add(normalized_phone)

        contacts.append(
            {
                "name": str(name).strip() if name is not None and str(name).strip() else None,
                "phone": normalized_phone,
                "email": str(email).strip() if email is not None and str(email).strip() else None,
                "tags": parse_tag_values(tag_value),
            }
        )
    return contacts


def get_or_create_tag(api: SupabaseRest, user_id: str, name: str, color: str) -> str:
    existing = api.request(
        "GET",
        "tags",
        query={
            "select": "id",
            "user_id": f"eq.{user_id}",
            "name": f"eq.{name}",
            "limit": "1",
        },
    )
    if existing:
        return existing[0]["id"]

    created = api.request(
        "POST",
        "tags",
        query={"select": "id"},
        body={"user_id": user_id, "name": name, "color": color},
        extra_headers={"Prefer": "return=representation"},
    )
    return created[0]["id"]


def fetch_existing_contacts(api: SupabaseRest, user_id: str, phones: list[str]) -> dict[str, dict[str, Any]]:
    if not phones:
        return {}
    resolved: dict[str, dict[str, Any]] = {}
    for phone_chunk in chunked(phones, 200):
        values = ",".join(phone_chunk)
        rows = api.request(
            "GET",
            "contacts",
            query={
                "select": "id,phone,name,email",
                "user_id": f"eq.{user_id}",
                "phone": f"in.({values})",
            },
        )
        for row in rows or []:
            resolved[row["phone"]] = row
    return resolved


def insert_missing_contacts(
    api: SupabaseRest,
    user_id: str,
    contacts: list[dict[str, str | None]],
    existing_by_phone: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    missing = [
        {
            "user_id": user_id,
            "phone": contact["phone"],
            "name": contact["name"],
            "email": contact["email"],
        }
        for contact in contacts
        if contact["phone"] not in existing_by_phone
    ]
    if not missing:
        return existing_by_phone

    inserted = api.request(
        "POST",
        "contacts",
        query={"select": "id,phone,name,email"},
        body=missing,
        extra_headers={"Prefer": "return=representation"},
    )
    merged = dict(existing_by_phone)
    for row in inserted or []:
        merged[row["phone"]] = row
    return merged


def existing_contact_tag_pairs(api: SupabaseRest, contact_ids: list[str], tag_ids: list[str]) -> set[tuple[str, str]]:
    if not contact_ids or not tag_ids:
        return set()
    tag_values = ",".join(tag_ids)
    pairs: set[tuple[str, str]] = set()
    for contact_chunk in chunked(contact_ids, 200):
        contact_values = ",".join(contact_chunk)
        rows = api.request(
            "GET",
            "contact_tags",
            query={
                "select": "contact_id,tag_id",
                "contact_id": f"in.({contact_values})",
                "tag_id": f"in.({tag_values})",
            },
        )
        pairs.update((row["contact_id"], row["tag_id"]) for row in rows or [])
    return pairs


def insert_contact_tags(api: SupabaseRest, rows: list[dict[str, str]]) -> None:
    if not rows:
        return
    for row_chunk in chunked(rows, 500):
        api.request(
            "POST",
            "contact_tags",
            body=row_chunk,
            extra_headers={"Prefer": "return=minimal"},
        )


def tag_color_for_name(name: str) -> str:
    palette = [
        "#3b82f6",
        "#10b981",
        "#f59e0b",
        "#ef4444",
        "#8b5cf6",
        "#06b6d4",
        "#84cc16",
        "#f97316",
    ]
    return palette[sum(ord(ch) for ch in name) % len(palette)]


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import campaign contacts from Excel and tag them in Supabase."
    )
    parser.add_argument(
        "--xlsx",
        default=str(ROOT / "campaign-2-30-05-2026.xlsx"),
        help="Path to the XLSX file to import.",
    )
    parser.add_argument("--user-id", required=True, help="Target auth.users/profile user_id UUID.")
    parser.add_argument(
        "--campaign-tag",
        default="campaign-2-30-05-2026",
        help="Campaign tag name to add to imported contacts.",
    )
    parser.add_argument(
        "--customer-tag",
        default="CUSTOMER",
        help="Customer tag name to add to imported contacts.",
    )
    args = parser.parse_args()

    supabase_url = resolve_env("NEXT_PUBLIC_SUPABASE_URL")
    service_role_key = resolve_env("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not service_role_key:
        raise RuntimeError(
            "Missing NEXT_PUBLIC_SUPABASE_URL or SUPABASE_SERVICE_ROLE_KEY in environment/.env."
        )

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        raise RuntimeError(f"Workbook not found: {xlsx_path}")

    contacts = read_contacts(xlsx_path)
    if not contacts:
        raise RuntimeError("No contacts found in workbook.")

    api = SupabaseRest(supabase_url.rstrip("/"), service_role_key)
    customer_tag_id = get_or_create_tag(api, args.user_id, args.customer_tag, "#3b82f6")
    campaign_tag_id = get_or_create_tag(api, args.user_id, args.campaign_tag, "#10b981")

    tag_ids_by_name: dict[str, str] = {
        args.customer_tag.casefold(): customer_tag_id,
        args.campaign_tag.casefold(): campaign_tag_id,
    }
    for contact in contacts:
        for tag_name in contact["tags"]:
            key = str(tag_name).casefold()
            if key in tag_ids_by_name:
                continue
            tag_ids_by_name[key] = get_or_create_tag(
                api,
                args.user_id,
                str(tag_name),
                tag_color_for_name(str(tag_name)),
            )

    existing_by_phone = fetch_existing_contacts(
        api, args.user_id, [str(contact["phone"]) for contact in contacts]
    )
    all_contacts_by_phone = insert_missing_contacts(api, args.user_id, contacts, existing_by_phone)

    contact_ids = [row["id"] for row in all_contacts_by_phone.values()]
    all_tag_ids = list(tag_ids_by_name.values())
    existing_pairs = existing_contact_tag_pairs(api, contact_ids, all_tag_ids)

    tag_rows: list[dict[str, str]] = []
    contacts_by_phone = {str(contact["phone"]): contact for contact in contacts}
    for phone, row in all_contacts_by_phone.items():
        contact = contacts_by_phone.get(phone, {})
        tag_names = [args.customer_tag, args.campaign_tag, *(contact.get("tags") or [])]
        for tag_name in tag_names:
            tag_id = tag_ids_by_name[str(tag_name).casefold()]
            pair = (row["id"], tag_id)
            if pair in existing_pairs:
                continue
            tag_rows.append({"contact_id": row["id"], "tag_id": tag_id})

    insert_contact_tags(api, tag_rows)

    print(f"Imported/checked {len(contacts)} contacts from {xlsx_path.name}")
    print(f"Resolved contacts: {len(all_contacts_by_phone)}")
    print(f"Added contact_tags rows: {len(tag_rows)}")
    print(f"Customer tag: {args.customer_tag} ({customer_tag_id})")
    print(f"Campaign tag: {args.campaign_tag} ({campaign_tag_id})")


if __name__ == "__main__":
    main()
